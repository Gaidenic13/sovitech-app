"""Generator of the synthetic XLSX companions of the IFC fixtures (ifc-input 5.3; prompt 3, 8).

    <python> -I -B fixtures/generators/xlsx/companions.py --out <folder>

writes two workbooks under <folder>/fixtures/xlsx/ and one ground-truth JSON per workbook under
<folder>/fixtures/xlsx/ground-truth/:

- tabel-camere.xlsx: room schedule. 17 guest rooms against the model's 16 (G4-9), all spaces
  (G9-6), areas with number formats, a total whose formula carries its cached value.
- lista-echipamente.xlsx: equipment schedule with tags (G4-3, G4-4, G8-5, G1-6, G1-1), "1.500 kW"
  as text with no locale (G8-3), an English-format sheet (G8-2, G8-3).

Determinism: openpyxl with fixed created and modified properties (its own save sets modified to
the clock, so the workbook is written through ExcelWriter directly), then the ZIP is rewritten
with stored entries and a fixed date_time (fixturelib/zip_tools.py). openpyxl writes formulas
without a cached value; the rewrite adds the value, computed here with Decimal, so a reader with
data_only=True sees what Excel would show.

Every workbook carries the TEST marking in its first row and in its document properties; the
building and every value are fictitious. Both Romanian diacritic forms appear on purpose:
comma-below (ș, ț) and cedilla (ş, ţ).
"""

from __future__ import annotations

import sys
from pathlib import Path

GENERATORS = Path(__file__).resolve().parents[1]
if str(GENERATORS) not in sys.path:
    sys.path.insert(0, str(GENERATORS))

import datetime  # noqa: E402
import io  # noqa: E402
import re  # noqa: E402
import zipfile  # noqa: E402
from decimal import Decimal  # noqa: E402
from typing import Any  # noqa: E402

from fixturelib.common import (  # noqa: E402
    AUTHOR,
    ORGANISATION_RO,
    TEST_MARK_EN,
    TEST_MARK_RO,
    parse_out,
    write_bytes,
    write_json,
)
from fixturelib.zip_tools import repack  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.writer.excel import ExcelWriter  # noqa: E402

FIXED_TIME = datetime.datetime(2026, 1, 15, 9, 0, 0)
GROUND_TRUTH_ABOUT = (
    "TEST ground truth of a synthetic XLSX fixture (fictitious building, not a real project), "
    "written by fixtures/generators/xlsx/companions.py with the workbook. Each cell is given with "
    "its cached value (what a reader with data_only=True sees), its number format and, for "
    "formulas, the formula text. Case ids name what a row is for; they are not app outputs."
)
TEST_BANNER = "DOCUMENT DE TEST: document sintetic, clădire fictivă. Nu este un proiect real."


def _workbook(title: str, subject: str) -> Workbook:
    wb = Workbook()
    props = wb.properties
    props.creator = ORGANISATION_RO
    props.lastModifiedBy = AUTHOR
    props.title = title
    props.subject = subject
    props.description = f"{TEST_MARK_EN} {TEST_MARK_RO}"
    props.keywords = "TEST FIXTURE, synthetic, fictitious building"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    return wb


def _save(wb: Workbook, cached: dict[tuple[str, str], str]) -> bytes:
    """Writes the workbook, then repacks it with fixed dates and the formulas' cached values."""
    buffer = io.BytesIO()
    archive = zipfile.ZipFile(buffer, "w", zipfile.ZIP_STORED)
    ExcelWriter(wb, archive).save()
    sheet_files = {
        f"xl/worksheets/sheet{index + 1}.xml": ws.title for index, ws in enumerate(wb.worksheets)
    }

    def transform(name: str, data: bytes) -> bytes:
        if name == "docProps/app.xml":
            text = data.decode("utf-8")
            text = re.sub(
                r"<Application>[^<]*</Application>",
                "<Application>SOVITECH fixture generator (openpyxl)</Application>",
                text,
            )
            return text.encode("utf-8")
        sheet = sheet_files.get(name)
        if sheet is None:
            return data
        text = data.decode("utf-8")
        for (sheet_name, cell), value in cached.items():
            if sheet_name != sheet:
                continue
            pattern = re.compile(rf'(<c r="{cell}"[^>]*>)(<f>[^<]*</f>)(<v\s*/>|<v></v>)?(</c>)')
            text, count = pattern.subn(
                lambda m, v=value: f"{m.group(1)}{m.group(2)}<v>{v}</v>{m.group(4)}", text
            )
            if count != 1:
                raise ValueError(f"{sheet_name}!{cell}: formula cell not found once in {name}")
        return text.encode("utf-8")

    return repack(buffer.getvalue(), transform)


def _cell_record(ws: Any, ref: str, cached: dict[tuple[str, str], str]) -> dict[str, Any]:
    cell = ws[ref]
    record: dict[str, Any] = {"sheet": ws.title, "cell": ref, "numberFormat": cell.number_format}
    if cell.data_type == "f":
        record["formula"] = cell.value
        record["cachedValue"] = cached[(ws.title, ref)]
        record["type"] = "formula"
    elif isinstance(cell.value, str):
        record["value"] = cell.value
        record["type"] = "text"
    elif cell.value is None:
        record["value"] = None
        record["type"] = "empty"
    else:
        record["value"] = str(cell.value)
        record["type"] = "number"
    return record


def _decimal_text(value: Decimal) -> str:
    return format(value.normalize(), "f")


# ---------------------------------------------------------------------------


ROOMS = [
    ("Subsol 1", "S1-01", "Parcare", "parcare", "800"),
    ("Subsol 1", "S1-02", "Cameră tehnică ventilare", "spațiu tehnic", "80"),
    ("Subsol 1", "S1-03", "Centrală termică", "spațiu tehnic", "93.6"),
    ("Parter", "P-01", "Recepție", "spațiu public", "240"),
    ("Parter", "P-02", "Restaurant", "spațiu public", "336"),
    ("Parter", "P-03", "Hol", "circulație", "160"),
    *[("Etaj 1", f"10{n}", f"Cameră 10{n}", "cameră de oaspeți", "26.4") for n in range(1, 9)],
    ("Etaj 1", "E1-H", "Hol E1", "circulație", "57.6"),
    *[("Etaj 2", f"20{n}", f"Cameră 20{n}", "cameră de oaspeți", "26.4") for n in range(1, 10)],
    ("Etaj 2", "E2-H", "Hol E2", "circulație", "57.6"),
]


def room_schedule() -> tuple[str, bytes, dict[str, Any]]:
    path = "fixtures/xlsx/tabel-camere.xlsx"
    wb = _workbook("TEST FIXTURE: Tabel camere (sintetic)", "Tabel de camere, clădire fictivă")
    ws = wb.active
    ws.title = "Camere"
    ws["A1"] = TEST_BANNER
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:E1")
    headers = ["Nivel", "Nr.", "Denumire", "Categorie", "Suprafață utilă (m²)"]
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=column, value=header)
        cell.font = Font(bold=True)
    records = []
    first = 4
    for offset, (level, number, name, category, area) in enumerate(ROOMS):
        row = first + offset
        ws.cell(row=row, column=1, value=level)
        ws.cell(row=row, column=2, value=number)
        ws.cell(row=row, column=3, value=name)
        ws.cell(row=row, column=4, value=category)
        area_cell = ws.cell(row=row, column=5, value=Decimal(area) if "." in area else int(area))
        area_cell.number_format = "#,##0.00"
        records.append(
            {
                "row": row,
                "number": number,
                "name": name,
                "category": category,
                "area": {
                    "cell": f"E{row}",
                    "value": area,
                    "unit": "m²",
                    "basis": "usable",
                    "numberFormat": "#,##0.00",
                },
                "guestRoom": category == "cameră de oaspeți",
            }
        )
    last = first + len(ROOMS) - 1
    total_row = last + 1
    ws.cell(row=total_row, column=3, value="Total").font = Font(bold=True)
    total_cell = ws.cell(row=total_row, column=5, value=f"=SUM(E{first}:E{last})")
    total_cell.number_format = "#,##0.00"
    total = sum((Decimal(room[4]) for room in ROOMS), Decimal(0))
    cached = {("Camere", f"E{total_row}"): _decimal_text(total)}

    summary = wb.create_sheet("Sumar")
    summary["A1"] = TEST_BANNER
    summary["A1"].font = Font(bold=True)
    guest_rooms = len([room for room in ROOMS if room[3] == "cameră de oaspeți"])
    summary["A3"] = "Număr camere de oaspeţi"  # ţ with a cedilla (U+0163)
    summary["B3"] = guest_rooms
    summary["B3"].number_format = "0"
    summary["A4"] = "Număr total de spații"  # ț with a comma below (U+021B)
    summary["B4"] = len(ROOMS)
    summary["B4"].number_format = "0"
    summary["A5"] = "Suprafață utilă totală (m²)"
    summary["B5"] = f"=Camere!E{total_row}"
    summary["B5"].number_format = "#,##0.00"
    cached[("Sumar", "B5")] = _decimal_text(total)

    data = _save(wb, cached)
    truth = {
        "about": GROUND_TRUTH_ABOUT,
        "file": path,
        "testMarking": {"cell": "Camere!A1", "text": TEST_BANNER},
        "properties": {
            "creator": ORGANISATION_RO,
            "created": "2026-01-15T09:00:00Z",
            "modified": "2026-01-15T09:00:00Z",
        },
        "sheets": ["Camere", "Sumar"],
        "headerRow": {"sheet": "Camere", "row": 3, "cells": headers},
        "rooms": records,
        "counts": {
            "guestRooms": {
                "value": guest_rooms,
                "cell": "Sumar!B3",
                "rowsCounted": "Camere rows with category 'cameră de oaspeți'",
                "cases": ["G4-9"],
                "note": (
                    "The model has 16 guest rooms (Cameră 101-108 and 201-208); this schedule "
                    "lists 209 too: a counts conflict, zero tolerance"
                ),
            },
            "allSpaces": {
                "value": len(ROOMS),
                "cell": "Sumar!B4",
                "cases": ["G9-6"],
                "note": "all_spaces, never room controllers",
            },
        },
        "cells": [
            _cell_record(ws, f"E{total_row}", cached),
            _cell_record(summary, "B3", cached),
            _cell_record(summary, "B4", cached),
            _cell_record(summary, "B5", cached),
        ],
        "diacritics": {
            "commaBelow": ["Sumar!A4 (ț)", "Camere!C7 (ț)"],
            "cedilla": ["Sumar!A3 (ţ)"],
        },
    }
    return path, data, truth


EQUIPMENT = [
    # tag, name, kind, level, system, parameter, value, number format, unit, remark, cases
    (
        "CTA-01",
        "Centrală tratare aer",
        "CTA",
        "Subsol 1",
        "Ventilare CTA-01",
        "Debit aer",
        9000,
        "#,##0",
        "m³/h",
        None,
        ["G3-1", "G4-3"],
    ),
    (
        "CTA-01",
        "Centrală tratare aer",
        "CTA",
        "Subsol 1",
        "Ventilare CTA-01",
        "Putere ventilator",
        Decimal("7.5"),
        "0.0",
        "kW",
        None,
        ["G4-3"],
    ),
    (
        "CTA-02",
        "Centrală tratare aer",
        "CTA",
        "Subsol 1",
        None,
        "Debit aer",
        "cca. 8.000",
        "General",
        "m³/h",
        "valoare aproximativă",
        ["rule 8 approximate"],
    ),
    (
        "CTA-03",
        "Centrală tratare aer",
        "CTA",
        "Subsol 1",
        "Ventilare parcare",
        "Debit aer",
        None,
        "General",
        "m³/h",
        "în curs de stabilire",
        ["G1-1"],
    ),
    (
        "CH-01",
        "Chiller răcit cu apă",
        "Chiller",
        "Subsol 1",
        "Apă răcită",
        "Putere frigorifică",
        430,
        "0",
        "kW",
        None,
        ["G8-5"],
    ),
    (
        "CH-01",
        "Chiller răcit cu apă",
        "Chiller",
        "Subsol 1",
        "Apă răcită",
        "Putere electrică absorbită",
        135,
        "0",
        "kW",
        None,
        ["G8-5"],
    ),
    (
        "CH-02",
        "Chiller răcit cu aer",
        "Chiller",
        "Subsol 1",
        "Apă răcită",
        "Putere frigorifică",
        None,
        "0",
        "kW",
        "nefurnizat de producător",
        ["G1-1"],
    ),
    (
        "P1.1",
        "Pompă circulație",
        "Pompă",
        "Subsol 1",
        "Apă răcită",
        "Configurație",
        "1+1R",
        "General",
        None,
        None,
        ["G4-4"],
    ),
    (
        "P1.1",
        "Pompă circulație",
        "Pompă",
        "Subsol 1",
        "Apă răcită",
        "Putere electrică absorbită",
        Decimal("5.5"),
        "0.0",
        "kW",
        None,
        [],
    ),
    (
        "P1.2",
        "Pompă circulație",
        "Pompă",
        "Subsol 1",
        "Apă răcită",
        "Configurație",
        "1+1R",
        "General",
        None,
        None,
        ["G4-4"],
    ),
    (
        "P2",
        "Pompă circulație",
        "Pompă",
        "Subsol 1",
        "Încălzire",
        "Configurație",
        "pompă dublă",
        "General",
        None,
        None,
        ["G4-4"],
    ),
    (
        "P2",
        "Pompă circulație",
        "Pompă",
        "Subsol 1",
        "Încălzire",
        "Putere motor",
        "1.500 kW",
        "General",
        None,
        "unitate în text",
        ["G8-3"],
    ),
    (
        "VE-P1",
        "Ventilator parcare",
        "Ventilator",
        "Subsol 1",
        "Ventilare parcare; Desfumare parcare",
        "Funcție",
        "ventilare CO și desfumare",
        "General",
        None,
        None,
        ["G11-4"],
    ),
    *[
        (
            f"VCV-1.0{n}",
            "Ventiloconvector",
            "VCV",
            "Etaj 1",
            "Apă răcită",
            "Putere frigorifică",
            Decimal("2.1"),
            "0.0",
            "kW",
            None,
            ["G3-8"],
        )
        for n in range(1, 9)
    ],
    *[
        (
            f"VCV-2.0{n}",
            "Ventiloconvector",
            "VCV",
            "Etaj 2",
            "Apă răcită",
            "Putere frigorifică",
            Decimal("2.1"),
            "0.0",
            "kW",
            None,
            ["G3-8", "IFC-10 second version"],
        )
        for n in range(1, 9)
    ],
    (
        "TA-01",
        "Tablou automatizare",
        "TA",
        "Subsol 1",
        "Automatizare",
        "Interfață",
        "Compatibil BMS",
        "General",
        None,
        None,
        ["G1-6"],
    ),
    (
        "CDI-01",
        "Centrală detecţie incendiu",
        "CDI",
        "Parter",
        "Detecţie incendiu",
        "Funcție",
        "monitorizare stare și alarme",
        "General",
        None,
        None,
        ["G11-3"],
    ),
    (
        "CET-01",
        "Contor energie termică",
        "Contor",
        "Subsol 1",
        "Încălzire",
        "Tip",
        "contor de energie termică",
        "General",
        None,
        None,
        [],
    ),
    (
        "DG-CT-01",
        "Detector de gaz",
        "Detector",
        "Subsol 1",
        "Detecție gaz",
        "Funcție",
        "detecție gaz metan",
        "General",
        None,
        None,
        [],
    ),
]


def equipment_schedule() -> tuple[str, bytes, dict[str, Any]]:
    path = "fixtures/xlsx/lista-echipamente.xlsx"
    wb = _workbook(
        "TEST FIXTURE: Listă echipamente (sintetic)", "Listă de echipamente, clădire fictivă"
    )
    ws = wb.active
    ws.title = "Echipamente"
    ws["A1"] = TEST_BANNER
    ws["A1"].font = Font(bold=True)
    ws.merge_cells("A1:I1")
    headers = [
        "Tag",
        "Denumire",
        "Tip",
        "Nivel",
        "Sistem",
        "Parametru",
        "Valoare",
        "Unitate",
        "Observații",
    ]
    for column, header in enumerate(headers, start=1):
        ws.cell(row=3, column=column, value=header).font = Font(bold=True)
    rows = []
    for offset, (
        tag,
        name,
        kind,
        level,
        system,
        parameter,
        value,
        number_format,
        unit,
        remark,
        cases,
    ) in enumerate(EQUIPMENT):
        row = 4 + offset
        for column, content in enumerate((tag, name, kind, level, system, parameter), start=1):
            ws.cell(row=row, column=column, value=content)
        value_cell = ws.cell(row=row, column=7, value=value)
        value_cell.number_format = number_format
        ws.cell(row=row, column=8, value=unit)
        ws.cell(row=row, column=9, value=remark)
        if value is None:
            written: Any = None
            value_type = "empty"
        elif isinstance(value, str):
            written = value
            value_type = "text"
        else:
            written = str(value)
            value_type = "number"
        rows.append(
            {
                "row": row,
                "tag": tag,
                "name": name,
                "level": level,
                "system": system,
                "parameter": parameter,
                "value": {
                    "cell": f"G{row}",
                    "type": value_type,
                    "value": written,
                    "numberFormat": number_format,
                },
                "unit": unit,
                "remark": remark,
                "cases": cases,
            }
        )

    maker = wb.create_sheet("Date producător")
    maker["A1"] = TEST_BANNER
    maker["A1"].font = Font(bold=True)
    maker["A3"] = "Item"
    maker["B3"] = "Value"
    maker["C3"] = "Value (number)"
    for cell in ("A3", "B3", "C3"):
        maker[cell].font = Font(bold=True)
    english = [
        ("Model", "EXC-AHU-12", None, "General"),
        ("Airflow", "9,000 m³/h", 9000, "#,##0"),
        ("Fan power", "7.5 kW", Decimal("7.5"), "0.0"),
        ("Supply air temperature", "18.0 °C", Decimal("18"), "0.0"),
    ]
    maker_rows = []
    for offset, (item, text, number, number_format) in enumerate(english):
        row = 4 + offset
        maker.cell(row=row, column=1, value=item)
        maker.cell(row=row, column=2, value=text)
        numeric = maker.cell(row=row, column=3, value=number)
        numeric.number_format = number_format
        maker_rows.append(
            {
                "row": row,
                "item": item,
                "text": {"cell": f"B{row}", "value": text, "locale": "en"},
                "number": None
                if number is None
                else {"cell": f"C{row}", "value": str(number), "numberFormat": number_format},
                "tag": "CTA-01",
                "cases": ["G8-2", "G8-3"],
            }
        )
    data = _save(wb, {})
    truth = {
        "about": GROUND_TRUTH_ABOUT,
        "file": path,
        "testMarking": {"cell": "Echipamente!A1", "text": TEST_BANNER},
        "properties": {
            "creator": ORGANISATION_RO,
            "created": "2026-01-15T09:00:00Z",
            "modified": "2026-01-15T09:00:00Z",
        },
        "sheets": ["Echipamente", "Date producător"],
        "headerRow": {"sheet": "Echipamente", "row": 3, "cells": headers},
        "rows": rows,
        "manufacturerSheet": {
            "sheet": "Date producător",
            "locale": "en",
            "note": (
                "English number format in text cells, numbers with number formats beside them "
                "(rule 8: locale per table)"
            ),
            "rows": maker_rows,
        },
        "diacritics": {
            "commaBelow": ["Echipamente!I3 header Observații (ț)"],
            "cedilla": ["Echipamente CDI-01 Denumire and Sistem (ţ)"],
        },
    }
    return path, data, truth


WORKBOOKS = (room_schedule, equipment_schedule)


def ground_truth_path(xlsx_path: str) -> str:
    name = xlsx_path.rsplit("/", 1)[-1].removesuffix(".xlsx")
    return f"fixtures/xlsx/ground-truth/{name}.json"


def main(argv: list[str] | None = None) -> None:
    args = parse_out(argv, __doc__ or "")
    for build in WORKBOOKS:
        path, data, truth = build()
        write_bytes(args.out, path, data)
        write_json(args.out, ground_truth_path(path), truth)


if __name__ == "__main__":
    main()
