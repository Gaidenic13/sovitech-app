"""A workbook's status says what its sheet coverage says, and a file is read only as the format it
was uploaded as (the phase 2 review).

- A workbook with a sheet not read (over the cell limit, a chart sheet, a read error) is partly
  analysed, counted in sheets: never "analysed" (rule 12, "Partial processing is shown with its
  coverage"; 2.3 ``partly_analysed``). The owner-facing line for sheets waits for the approver
  (P-2-XLSX-SHEETS); the stored status is data, not new wording. With no sheet read, its analysis
  failed, and each sheet's reason stays in the coverage for the engineer.
- A workbook uploaded as a PDF is not read at all: its analysis fails with ``format_mismatch``, so
  no XLSX text, sheet coverage or status is stored under a PDF record (rule 12, 2.3). A workbook
  uploaded as a ZIP archive is stored whole, as an archive (R-014).

Every workbook is synthetic TEST content made in the test with openpyxl, in a temporary folder.

Ids: F-EXTRACT-01, F-INGEST-03, F-INGEST-05, R-014, R-015, G12-4 (extractor half).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.chart import BarChart, Reference

from sovitech_extractor import contract
from sovitech_extractor.extract import file_hash, run_job

PROJECT = "0192f0a0-0000-7000-8000-00000000e201"
DOCUMENT = "0192f0a0-0000-7000-8000-00000000e202"


def _request(path: Path, declared: str, max_cells: int = 1000) -> contract.ExtractionRequest:
    value: dict[str, Any] = {
        "contractVersion": "1.0.0",
        "job": {"projectId": PROJECT, "documentId": DOCUMENT, "contentHash": file_hash(path)},
        "declaredFormat": declared,
        "ifcValues": False,
        "datasets": [],
        "derivatives": [],
        "limits": {"maxPages": 100, "maxCellsPerSheet": max_cells, "wallClockSeconds": 60},
    }
    return contract.parse_request(value)


def _workbook(path: Path, *, big_rows: int = 10, chart: bool = False) -> Path:
    workbook = openpyxl.Workbook()
    small = workbook.active
    small.title = "TEST small"
    small["A1"] = "TEST camera"
    small["B1"] = 1
    big = workbook.create_sheet("TEST big")
    for row in range(1, big_rows + 1):
        big.cell(row=row, column=1, value=f"TEST row {row}")
        big.cell(row=row, column=2, value=row)
    if chart:
        sheet = workbook.create_chartsheet("TEST chart")
        bars = BarChart()
        bars.add_data(Reference(big, min_col=2, min_row=1, max_row=3))
        sheet.add_chart(bars)
    workbook.save(path)
    return path


def _checked(request: contract.ExtractionRequest, path: Path) -> contract.ExtractionOutput:
    output = run_job(request, path)
    # The writer refuses what the API would refuse, and the output answers its request.
    contract.dump_output(output)
    assert contract.output_answers_request(request, output) == ()
    return output


def test_G12_4_a_sheet_over_the_cell_limit_makes_the_workbook_partly_analysed_in_sheets(
    tmp_path: Path,
) -> None:
    path = _workbook(tmp_path / "limit.xlsx", big_rows=10)
    output = _checked(_request(path, "xlsx", max_cells=4), path)
    assert output.analysis == contract.PartlyAnalysedStatus("partly_analysed", "sheets", 1, 2)
    assert output.coverage.sheets is not None
    assert [(entry.sheet, entry.status, entry.reason) for entry in output.coverage.sheets] == [
        ("TEST small", "read", None),
        ("TEST big", "not_read", "cell_limit"),
    ]
    # Within the limit, every sheet is read, and the workbook is analysed.
    whole = _checked(_request(path, "xlsx", max_cells=1000), path)
    assert whole.analysis == contract.AnalysedStatus("analysed")


def test_f_ingest_05_a_chart_sheet_not_read_makes_the_workbook_partly_analysed(
    tmp_path: Path,
) -> None:
    path = _workbook(tmp_path / "chart.xlsx", chart=True)
    output = _checked(_request(path, "xlsx"), path)
    assert output.analysis == contract.PartlyAnalysedStatus("partly_analysed", "sheets", 2, 3)
    assert output.coverage.sheets is not None
    assert [entry.reason for entry in output.coverage.sheets if entry.status == "not_read"] == [
        "chart_sheet"
    ]


def test_f_ingest_05_a_workbook_with_no_sheet_read_failed_and_keeps_its_sheet_coverage(
    tmp_path: Path,
) -> None:
    path = _workbook(tmp_path / "nothing.xlsx", big_rows=10)
    output = _checked(_request(path, "xlsx", max_cells=1), path)
    assert output.analysis == contract.AnalysisFailedStatus("failed", "unreadable_file")
    assert output.xlsx is None
    assert output.coverage.sheets is not None
    assert {entry.status for entry in output.coverage.sheets} == {"not_read"}


def test_f_ingest_03_a_workbook_uploaded_as_a_pdf_is_not_read_and_fails_with_format_mismatch(
    tmp_path: Path,
) -> None:
    path = _workbook(tmp_path / "renamed.pdf")
    output = _checked(_request(path, "pdf"), path)
    assert (output.format, output.analysis) == (
        "other",
        contract.AnalysisFailedStatus("failed", "format_mismatch"),
    )
    assert (output.pdf, output.xlsx, output.coverage) == (None, None, contract.Coverage())
    # The contract refuses a workbook's output under a PDF request, in both languages.
    as_workbook = run_job(_request(path, "xlsx"), path)
    problems = contract.output_answers_request(_request(path, "pdf"), as_workbook)
    assert [problem.path for problem in problems] == ["/format"]


def test_r_014_a_workbook_uploaded_as_a_zip_archive_is_stored_whole(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "archive.zip")
    output = _checked(_request(path, "zip"), path)
    assert (output.format, output.analysis) == (
        "zip",
        contract.StoredOnlyStatus("stored_only", "ZIP archive"),
    )
    assert output.xlsx is None
